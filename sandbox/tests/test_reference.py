"""The hand-written reference pipeline must find every planted exception for all 12 clients."""

from __future__ import annotations

import hashlib
import json
import sys
from decimal import Decimal
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).parent))
from reference_recon import run  # noqa: E402

from tieout_lib import validate_result  # noqa: E402

DATA = Path(__file__).resolve().parents[2] / "data" / "demo"
CLIENTS = sorted(p.name for p in DATA.iterdir() if (p / "expected.json").exists())


def key(e: dict) -> tuple:
    return (e["kind"], Decimal(e["amount"]), e.get("bank_ref") or None, e.get("gl_ref") or None)


@pytest.mark.parametrize("client", CLIENTS)
def test_reference_finds_every_planted_exception(client: str, tmp_path: Path) -> None:
    exp = json.loads((DATA / client / "expected.json").read_text())
    out = run(DATA / client, tmp_path)
    res, errors = validate_result((tmp_path / "result.json").read_text())
    assert errors == [], errors
    assert res.difference == Decimal("0.00")
    assert f"{res.bank_ending_balance:.2f}" == exp["bank_ending_balance"]
    assert f"{res.gl_ending_balance:.2f}" == exp["gl_ending_balance"]
    assert f"{res.adjusted_bank_balance:.2f}" == exp["adjusted_bank_balance"]
    assert f"{res.adjusted_book_balance:.2f}" == exp["adjusted_book_balance"]
    found = sorted(key(e.model_dump(mode="json")) for e in res.exceptions)
    want = sorted(key(e) for e in exp["exceptions"])
    assert found == want
    got_ajes = sorted((a.debit_account, a.credit_account, f"{a.amount:.2f}") for a in res.ajes)
    want_ajes = sorted((a["debit_account"], a["credit_account"], a["amount"]) for a in exp["ajes"])
    assert got_ajes == want_ajes
    assert res.status == exp["status"]
    assert res.bank_line_count == exp["bank_line_count"]


def test_outputs_are_byte_identical(tmp_path: Path) -> None:
    client = DATA / "blue-harbor-coffee"
    h1 = run(client, tmp_path / "a")["hashes"]
    h2 = run(client, tmp_path / "b")["hashes"]
    assert h1 == h2
    for name, digest in h1.items():
        assert hashlib.sha256((tmp_path / "a" / name).read_bytes()).hexdigest() == digest


def test_workpaper_has_formulas(tmp_path: Path) -> None:
    from openpyxl import load_workbook

    run(DATA / "blue-harbor-coffee", tmp_path)
    wb = load_workbook(tmp_path / "workpaper.xlsx")
    assert wb.sheetnames == ["Summary", "Exceptions", "Matches", "Proposed AJEs", "Inputs"]
    ws = wb["Summary"]
    assert str(ws["B10"].value).startswith("=")
    assert str(ws["B17"].value).startswith("=ROUND(")
